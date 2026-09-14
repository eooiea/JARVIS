"""
Scenario 1 Test Data Generator
Generates mock order files for Coupang, Naver Smartstore, and Gmarket.
"""
from pathlib import Path
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

def apply_table_styles(ws, header_fill_hex: str = "34495E"):
    header_font = Font(name="Pretendard", size=11, bold=True, color="FFFFFF")
    header_fill = PatternFill(start_color=header_fill_hex, end_color=header_fill_hex, fill_type="solid")
    data_font = Font(name="Pretendard", size=10)
    
    thin = Side(border_style="thin", color="D0D5DD")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    center_align = Alignment(horizontal="center", vertical="center")
    left_align = Alignment(horizontal="left", vertical="center")
    
    # Format headers
    for col_idx in range(1, ws.max_column + 1):
        cell = ws.cell(row=1, column=col_idx)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = center_align
        cell.border = border

    # Format data
    for row in range(2, ws.max_row + 1):
        for col in range(1, ws.max_column + 1):
            cell = ws.cell(row=row, column=col)
            cell.font = data_font
            cell.border = border
            if isinstance(cell.value, int):
                cell.alignment = center_align
            elif col in [1, 2, 4, 5]:  # IDs, phones
                cell.alignment = center_align
            else:
                cell.alignment = left_align

    # Auto-adjust column widths
    for col in ws.columns:
        col_letter = get_column_letter(col[0].column)
        max_len = max(len(str(cell.value or '')) for cell in col)
        ws.column_dimensions[col_letter].width = max(max_len * 1.6, 12)

def generate_test_data():
    orders_dir = Path("c:/Coding/jarvis/orders")
    orders_dir.mkdir(parents=True, exist_ok=True)
    
    # 1. 쿠팡 (Coupang)
    wb_coupang = openpyxl.Workbook()
    ws_coupang = wb_coupang.active
    ws_coupang.title = "쿠팡주문목록"
    
    coupang_headers = [
        "묶음배송번호", "주문번호", "받는사람", "수하인전화번호", "구매자전화번호",
        "수하인주소", "등록상품명", "구매수량(개)", "배송메세지", "결제일시"
    ]
    ws_coupang.append(coupang_headers)
    
    coupang_rows = [
        ["CP-20260914-001", "202609141001", "김철수", "010-1111-2222", "010-1111-2222", "서울시 강남구 테헤란로 123 4층", "프리미엄 캠핑의자 블랙", 2, "문 앞에 놓아주세요", "2026-09-14 08:30:12"],
        ["CP-20260914-002", "202609141002", "이영희", "010-3333-4444", "010-3333-4444", "경기도 성남시 분당구 판교역로 100 101동 502호", "캠핑 롤테이블 우드", 1, "배송 전 연락바랍니다", "2026-09-14 09:12:45"],
        ["CP-20260914-003", "202609141003", "박민수", "010-5555-6666", "010-5555-6666", "부산시 해운대구 마린시티2로 33 2001호", "초경량 침낭 올리브", 1, "경비실에 맡겨주세요", "2026-09-14 09:40:22"],
        ["CP-20260914-004", "202609141004", "정하나", "01077778888", "01077778888", "대구시 수성구 범어천로 50 103동 1204호", "캠핑 감성 랜턴", 3, "파손주의 부탁드립니다", "2026-09-14 09:55:01"],
    ]
    for r in coupang_rows:
        ws_coupang.append(r)
    apply_table_styles(ws_coupang, header_fill_hex="C0392B")  # Red-ish for Coupang
    wb_coupang.save(orders_dir / "쿠팡_주문내역_20260914.xlsx")

    # 2. 스마트스토어 (Naver Smartstore)
    wb_naver = openpyxl.Workbook()
    ws_naver = wb_naver.active
    ws_naver.title = "발주발송관리"
    
    naver_headers = [
        "상품주문번호", "주문번호", "수취인명", "연락처1", "연락처2",
        "기본배송지", "상세배송지", "상품명", "옵션정보", "수량", "배송메모", "주문일시"
    ]
    ws_naver.append(naver_headers)
    
    naver_rows = [
        ["NV20260914001", "20260914-NV01", "박민수", "010.5555.6666", "051-123-4567", "부산시 해운대구 마린시티2로 33", "2001호", "초경량 침낭 올리브", "색상: 올리브", 2, "경비실에 맡겨주세요", "2026-09-14 09:42:10"],
        ["NV20260914002", "20260914-NV02", "최유진", "010-9999-0000", "", "인천시 연수구 송도과학로 16", "1502호", "감성 인디언 행어", "기본형", 1, "문 앞", "2026-09-14 10:01:00"],
        ["NV20260914003", "20260914-NV03", "최유진", "010-9999-0000", "", "인천시 연수구 송도과학로 16", "1502호", "감성 인디언 행어", "추가 고리 5P", 2, "문 앞", "2026-09-14 10:03:25"],
        ["NV20260914004", "20260914-NV04", "한동훈", "010-2468-1357", "", "대전시 유성구 대학로 99", "공학관 301호", "폴딩 캠핑 박스 50L", "블랙", 4, "부재 시 문 앞", "2026-09-14 10:15:40"],
    ]
    for r in naver_rows:
        ws_naver.append(r)
    apply_table_styles(ws_naver, header_fill_hex="27AE60")  # Green for Naver
    wb_naver.save(orders_dir / "스마트스토어_발주발송관리_20260914.xlsx")

    # 3. 지마켓 (Gmarket)
    wb_gmarket = openpyxl.Workbook()
    ws_gmarket = wb_gmarket.active
    ws_gmarket.title = "주문배송조회"
    
    gmarket_headers = [
        "주문번호", "성명", "핸드폰", "전화번호", "주소", "주문상품명", "수량", "배송요구사항", "주문일자"
    ]
    ws_gmarket.append(gmarket_headers)
    
    gmarket_rows = [
        ["GM-20260914-11", "김철수", "010-1111-2222", "02-987-6543", "서울시 강남구 테헤란로 123 4층", "프리미엄 캠핑의자 블랙", 1, "문 앞에 놓아주세요", "2026-09-14 09:05:00"],
        ["GM-20260914-12", "강태풍", "010-8888-9999", "", "광주시 서구 상무중앙로 45 301호", "야외용 방수 타프 3x3", 1, "배송 전 문자 요망", "2026-09-14 10:05:12"],
        ["GM-20260914-13", "조미래", "010-3456-7890", "", "세종특별자치시 한누리대로 200 (도담동, 도램마을 1단지)", "휴대용 가스 버너 버너팩", 1, "안전배송", "2026-09-14 10:20:00"],
    ]
    for r in gmarket_rows:
        ws_gmarket.append(r)
    apply_table_styles(ws_gmarket, header_fill_hex="2980B9")  # Blue for Gmarket
    wb_gmarket.save(orders_dir / "지마켓_주문배송조회_20260914.xlsx")

    print(f"Generated 3 order Excel files successfully in {orders_dir}")

if __name__ == "__main__":
    generate_test_data()
