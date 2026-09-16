"""Generate 10 Diverse Scenario 1 Test Sets in tests/data/scenario1/."""

from pathlib import Path
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter


def apply_table_styles(ws, header_fill_hex: str = "34495E"):
    header_font = Font(name="Pretendard", size=10, bold=True, color="FFFFFF")
    header_fill = PatternFill(start_color=header_fill_hex, end_color=header_fill_hex, fill_type="solid")
    data_font = Font(name="Pretendard", size=9)
    
    thin = Side(border_style="thin", color="D0D5DD")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    center_align = Alignment(horizontal="center", vertical="center")
    left_align = Alignment(horizontal="left", vertical="center")
    
    for col_idx in range(1, ws.max_column + 1):
        cell = ws.cell(row=1, column=col_idx)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = center_align
        cell.border = border

    for row in range(2, ws.max_row + 1):
        for col in range(1, ws.max_column + 1):
            cell = ws.cell(row=row, column=col)
            cell.font = data_font
            cell.border = border
            if isinstance(cell.value, int):
                cell.alignment = center_align
            else:
                cell.alignment = left_align

    for col in ws.columns:
        col_letter = get_column_letter(col[0].column)
        max_len = max(len(str(cell.value or '')) for cell in col)
        ws.column_dimensions[col_letter].width = max(max_len * 1.5, 10)


def create_simple_excel(file_path: Path, sheet_title: str, headers: list, rows: list, color: str = "2C3E50"):
    file_path.parent.mkdir(parents=True, exist_ok=True)
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = sheet_title
    ws.append(headers)
    for r in rows:
        ws.append(r)
    apply_table_styles(ws, header_fill_hex=color)
    wb.save(file_path)


def generate_all_10_testsets(base_dir: Path):
    print(f"Generating 10 test sets in {base_dir}...")

    # Set 01: 쿠팡 + 스마트스토어 + 지마켓 -> CJ대한통운 (기본형 11건)
    s1_dir = base_dir / "set01"
    create_simple_excel(
        s1_dir / "orders" / "쿠팡_주문.xlsx", "쿠팡",
        ["묶음번호", "받는사람", "수하인전화번호", "수하인주소", "등록상품명", "구매수량(개)", "배송메세지"],
        [
            ["CP-01", "김철수", "010-1111-2222", "서울시 강남구 테헤란로 123", "캠핑의자", 2, "문 앞"],
            ["CP-02", "이영희", "010.3333.4444", "경기도 성남시 분당구 판교역로 100", "롤테이블", "1개", "배송 전 연락"],
            ["CP-03", "박민수", "01055556666", "부산시 해운대구 마린시티2로 33", "초경량 침낭", 1, "경비실"],
            ["CP-04", "정하나", "010-7777-8888", "대구시 수성구 범어천로 50", "감성 랜턴", 3, "파손주의"],
        ], "C0392B"
    )
    create_simple_excel(
        s1_dir / "orders" / "스마트스토어_주문.xlsx", "스마트스토어",
        ["주문번호", "수취인명", "연락처1", "기본배송지", "상세배송지", "상품명", "수량", "배송메모"],
        [
            ["NV-01", "박민수", "010.5555.6666", "부산시 해운대구 마린시티2로 33", "2001호", "초경량 침낭", 2, "경비실"],
            ["NV-02", "최유진", "010-9999-0000", "인천시 연수구 송도과학로 16", "1502호", "인디언 행어", 1, "문 앞"],
            ["NV-03", "한동훈", "010-2468-1357", "대전시 유성구 대학로 99", "301호", "캠핑 박스", 4, ""],
            ["NV-04", "강감찬", "010-3333-7777", "서울시 관악구 낙성대로 10", "2층", "캠핑 매트", 1, "벨 금지"],
        ], "27AE60"
    )
    create_simple_excel(
        s1_dir / "orders" / "지마켓_주문.xlsx", "지마켓",
        ["주문번호", "성명", "핸드폰", "주소", "주문상품명", "수량", "배송요구사항"],
        [
            ["GM-01", "김철수", "010-1111-2222", "서울시 강남구 테헤란로 123", "캠핑의자", 1, "문 앞"],
            ["GM-02", "강태풍", "010-8888-9999", "광주시 서구 상무중앙로 45", "방수 타프", 1, "문자 요망"],
            ["GM-03", "조미래", "010-3456-7890", "세종시 한누리대로 200", "휴대용 가스 버너", 1, "안전배송"],
        ], "2980B9"
    )

    # Set 02: 스마트스토어 + 11번가 -> 우체국택배
    s2_dir = base_dir / "set02"
    create_simple_excel(
        s2_dir / "orders" / "스마트스토어_발주.xlsx", "스마트스토어",
        ["수취인명", "연락처1", "기본배송지", "상세배송지", "상품명", "수량", "배송메모"],
        [
            ["오세훈", "01011112222", "서울시 중구 세종대로 110", "본관", "스마트 텀블러", 2, "택배함"],
            ["박형준", "010-2222-3333", "부산시 연제구 중앙대로 1001", "시청사", "무선 가습기", 1, "부재시 연락"],
        ], "27AE60"
    )
    create_simple_excel(
        s2_dir / "orders" / "11번가_주문.xlsx", "11번가",
        ["주문번호", "수령인", "휴대전화", "배송주소", "상품명", "주문수량", "배송메모"],
        [
            ["11ST-01", "유정복", "010.4444.5555", "인천시 남동구 정각로 29", "탁상용 선풍기", "2EA", "문앞 보관"],
            ["11ST-02", "김동연", "01066667777", "경기도 수원시 영통구 도청로 30", "모니터 거치대", 1, ""],
        ], "D35400"
    )

    # Set 03: 쿠팡 + 카카오선물하기 -> 한진택배
    s3_dir = base_dir / "set03"
    create_simple_excel(
        s3_dir / "orders" / "쿠팡_로켓.xlsx", "쿠팡",
        ["받는사람", "수하인전화번호", "수하인주소", "등록상품명", "구매수량(개)", "배송메세지"],
        [
            ["이재명", "010-5555-1234", "서울시 영등포구 의사당대로 1", "원두 커피 1kg", 3, "의원실"],
            ["윤석열", "010-7777-4321", "서울시 용산구 이태원로 22", "고급 머그컵", 2, "직접 전달"],
        ], "C0392B"
    )
    create_simple_excel(
        s3_dir / "orders" / "카카오선물하기_주문.xlsx", "카카오",
        ["선물주문번호", "받는분", "수신자연락처", "배송지", "상품명", "수량", "배송요청"],
        [
            ["KAKAO-01", "한동훈", "010-8888-2468", "서울시 종로구 삼청로 1", "홍삼 선물세트", 1, "경비실"],
        ], "F39C12"
    )

    # Set 04: 무신사 + 29CM -> 롯데택배
    s4_dir = base_dir / "set04"
    create_simple_excel(
        s4_dir / "orders" / "무신사_출고.xlsx", "무신사",
        ["주문고유번호", "수령자명", "핸드폰번호", "배송지주소", "상품명(옵션)", "출고수량", "배송요청메모"],
        [
            ["MS-01", "차은우", "010-9999-1111", "서울시 강남구 압구정로 10", "오버핏 후드 블랙 (L)", 1, "문 앞"],
            ["MS-02", "송혜교", "010.8888.2222", "서울시 용산구 한남대로 20", "와이드 데님 팬츠 (S)", 2, "배송전 연락"],
        ], "1B2631"
    )
    create_simple_excel(
        s4_dir / "orders" / "29CM_주문.xlsx", "29CM",
        ["주문번호", "받는분성함", "연락처", "수령주소", "아이템명", "수량", "요청사항"],
        [
            ["29CM-01", "공유", "01077773333", "서울시 마포구 독막로 30", "미니멀 울 자켓 (48)", 1, "부재시 경비실"],
        ], "2C3E50"
    )

    # Set 05: 지마켓 + 옥션 (ESM) -> 로젠택배
    s5_dir = base_dir / "set05"
    create_simple_excel(
        s5_dir / "orders" / "지마켓_ESM.xlsx", "지마켓",
        ["성명", "전화번호", "주소", "상품명", "수량", "배송메모"],
        [
            ["유재석", "010-1234-7890", "서울시 강남구 논현로 100", "비타민C 1000", 3, "경비실"],
        ], "2980B9"
    )
    create_simple_excel(
        s5_dir / "orders" / "옥션_ESM.xlsx", "옥션",
        ["구매자명", "휴대폰", "수취인주소", "품목", "수량", "메모"],
        [
            ["강호동", "010.9876.5432", "서울시 송파구 올림픽로 200", "루테인 오메가3", 2, "문 앞"],
            ["신동엽", "01054321234", "서울시 마포구 상암산로 50", "밀크씨슬 골드", 1, "안전배송"],
        ], "E74C3C"
    )

    # Set 06: 스마트스토어 + 쿠팡 -> 사용자 커스텀 템플릿
    s6_dir = base_dir / "set06"
    # 자사 커스텀 템플릿 파일 생성
    create_simple_excel(
        s6_dir / "templates" / "자사_송장양식.xlsx", "송장양식",
        ["고객명", "휴대폰번호", "전체주소", "출고상품", "출고수량", "배송특이사항"],
        [], "8E44AD"
    )
    create_simple_excel(
        s6_dir / "orders" / "스마트스토어.xlsx", "스마트스토어",
        ["수취인명", "연락처1", "기본배송지", "상세배송지", "상품명", "수량", "배송메모"],
        [
            ["손흥민", "010-7777-7777", "강원도 춘천시 동면 만천리 1", "707호", "축구공 프로", 1, "문 앞"],
        ], "27AE60"
    )
    create_simple_excel(
        s6_dir / "orders" / "쿠팡.xlsx", "쿠팡",
        ["받는사람", "수하인전화번호", "수하인주소", "등록상품명", "구매수량(개)", "배송메세지"],
        [
            ["이강인", "010-1919-1919", "인천시 남동구 간석동 20", "정강이 보호대", 2, "배송전 문자"],
        ], "C0392B"
    )

    # Set 07: 에이블리 + 지그재그 + 올리브영 -> CJ대한통운
    s7_dir = base_dir / "set07"
    create_simple_excel(
        s7_dir / "orders" / "에이블리.xlsx", "에이블리",
        ["수취인", "전화번호", "배송주소", "상품", "개수", "요청사항"],
        [["김민희", "010-1111-3333", "서울시 용산구 소월로 1", "립스틱 코랄", "2개", "문 앞"]], "9B59B6"
    )
    create_simple_excel(
        s7_dir / "orders" / "지그재그.xlsx", "지그재그",
        ["받는사람이름", "휴대폰", "주소지", "상품상세", "수량", "메모"],
        [["전지현", "010.2222.4444", "서울시 강남구 도산대로 10", "수분 크림", "1EA", "부재시 문앞"]], "E91E63"
    )
    create_simple_excel(
        s7_dir / "orders" / "올리브영.xlsx", "올리브영",
        ["수령자", "연락처", "배송지", "아이템", "수량", "배송요청"],
        [["아이유", "01033335555", "서울시 광진구 구의강변로 5", "선크림 50ml", 3, "경비실"]], "2ECC71"
    )

    # Set 08: 쿠팡 + 11번가 -> 우체국택배 (전화번호 노이즈 케이스)
    s8_dir = base_dir / "set08"
    create_simple_excel(
        s8_dir / "orders" / "쿠팡_안심번호.xlsx", "쿠팡",
        ["받는사람", "수하인전화번호", "수하인주소", "등록상품명", "구매수량(개)", "배송메세지"],
        [
            ["김연아", "050412345678", "서울시 송파구 백제고분로 10", "피겨 스케이트 날", 1, "문 앞"],
            ["박태환", "02-123-4567", "서울시 노원구 화랑로 20", "수영모", 2, "직접 수령"],
        ], "C0392B"
    )
    create_simple_excel(
        s8_dir / "orders" / "11번가_전화번호.xlsx", "11번가",
        ["수령인", "휴대전화", "배송주소", "상품명", "주문수량", "배송메모"],
        [
            ["장미란", "0314445555", "경기도 고양시 일산동구 중앙로 30", "리프팅 벨트", 1, "체육관"],
        ], "D35400"
    )

    # Set 09: 스마트스토어 단독 15건 대량 주문 -> 롯데택배
    s9_dir = base_dir / "set09"
    s9_rows = [
        [f"고객_{i:02d}", f"010-{i*1111:04d}-{i*2222:04d}", "서울시 마포구 월드컵북로", f"{100+i}호", f"블랙 티셔츠 ({i})", i % 3 + 1, "안전배송"]
        for i in range(1, 16)
    ]
    create_simple_excel(
        s9_dir / "orders" / "스마트스토어_대량발주.xlsx", "스마트스토어",
        ["수취인명", "연락처1", "기본배송지", "상세배송지", "상품명", "수량", "배송메모"],
        s9_rows, "27AE60"
    )

    # Set 10: 쿠팡 + 지마켓 + 스마트스토어 -> 로젠택배 (자가수정 케이스 검증용)
    s10_dir = base_dir / "set10"
    create_simple_excel(
        s10_dir / "orders" / "쿠팡_주문.xlsx", "쿠팡",
        ["받는사람", "수하인전화번호", "수하인주소", "등록상품명", "구매수량(개)", "배송메세지"],
        [["김고은", "010-5555-9999", "서울시 성동구 성수일로 10", "도깨비 캔들", 1, "문 앞"]], "C0392B"
    )
    create_simple_excel(
        s10_dir / "orders" / "지마켓_주문.xlsx", "지마켓",
        ["성명", "핸드폰", "주소", "주문상품명", "수량", "배송요구사항"],
        [["공지철", "010.7777.1111", "서울시 용산구 이태원로 50", "커피 머신", 1, "부재시 연락"]], "2980B9"
    )
    create_simple_excel(
        s10_dir / "orders" / "스마트스토어_주문.xlsx", "스마트스토어",
        ["수취인명", "연락처1", "기본배송지", "상세배송지", "상품명", "수량", "배송메모"],
        [["이동욱", "01099993333", "서울시 영등포구 여의대로 24", "101호", "블루투스 마이크", 2, "경비실"]], "27AE60"
    )

    print("Successfully generated all 10 scenario test sets!")


if __name__ == "__main__":
    generate_all_10_testsets(Path("tests/data/scenario1"))
