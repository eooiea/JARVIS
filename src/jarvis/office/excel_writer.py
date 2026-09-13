"""Smart styled Excel file generator using openpyxl."""

from pathlib import Path
from typing import Union, List, Dict, Any
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from jarvis.utils.logger import get_logger

logger = get_logger("office.excel_writer")


class ExcelWriter:
    """Creates beautifully styled Excel spreadsheets with headers, borders, and auto column widths."""

    def __init__(self):
        # Professional Navy/Blue theme
        self.header_font = Font(name="맑은 고딕", size=11, bold=True, color="FFFFFF")
        self.header_fill = PatternFill(fill_type="solid", start_color="1F4E78", end_color="1F4E78")
        self.body_font = Font(name="맑은 고딕", size=10)
        
        thin_border = Side(style="thin", color="D9D9D9")
        self.cell_border = Border(left=thin_border, right=thin_border, top=thin_border, bottom=thin_border)
        
        self.center_align = Alignment(horizontal="center", vertical="center")
        self.left_align = Alignment(horizontal="left", vertical="center")
        self.right_align = Alignment(horizontal="right", vertical="center")

    def _normalize_data_rows(self, data: Union[str, Dict[str, Any], List[Dict[str, Any]]]) -> List[Dict[str, Any]]:
        """Normalize various complex/nested JSON shapes into a flat list of row dictionaries."""
        if isinstance(data, str):
            import json
            try:
                data = json.loads(data)
            except Exception:
                return [{"내용": data}]

        if isinstance(data, list):
            # If list contains dicts, return them; if list contains primitives, wrap each
            return [row if isinstance(row, dict) else {"값": row} for row in data]
        
        elif isinstance(data, dict):
            # 1. Look for known common list keys or any list of dicts inside
            for key, val in data.items():
                if isinstance(val, list) and len(val) > 0:
                    # Found an array! Check if elements are dicts
                    if all(isinstance(elem, dict) for elem in val):
                        return val
                    else:
                        return [{"항목": elem} for elem in val]

            # 2. Check if the dict has child dicts (e.g. key-value map of objects)
            rows = []
            flat_dict = {}
            for k, v in data.items():
                if isinstance(v, dict):
                    # Child dict -> flatten with parent key prefix
                    for sub_k, sub_v in v.items():
                        flat_dict[f"{k}_{sub_k}"] = sub_v
                else:
                    flat_dict[k] = v
            return [flat_dict]
            
        return []

    def _format_cell_value(self, val: Any) -> Any:
        """Ensure cell value is a primitive accepted by openpyxl."""
        if val is None:
            return ""
        if isinstance(val, (int, float, bool)):
            return val
        if isinstance(val, (dict, list)):
            import json
            try:
                return json.dumps(val, ensure_ascii=False)
            except Exception:
                return str(val)
        return str(val)

    def create_styled_excel(
        self,
        data: Union[Dict[str, Any], List[Dict[str, Any]]],
        output_path: Union[Path, str],
        sheet_name: str = "Data",
    ) -> Path:
        """Create styled .xlsx file from structured data."""
        out_file = Path(output_path).resolve()
        out_file.parent.mkdir(parents=True, exist_ok=True)

        rows = self._normalize_data_rows(data)
        if not rows:
            rows = [{"내용": "데이터 없음"}]

        # Collect unique headers across all rows preserving order
        headers_set = set()
        headers = []
        for r in rows:
            for k in r.keys():
                if k not in headers_set:
                    headers_set.add(k)
                    headers.append(k)

        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = sheet_name

        # 1. Write Header Row
        for col_idx, header in enumerate(headers, start=1):
            cell = ws.cell(row=1, column=col_idx, value=str(header))
            cell.font = self.header_font
            cell.fill = self.header_fill
            cell.alignment = self.center_align
            cell.border = self.cell_border
        ws.row_dimensions[1].height = 26

        # 2. Write Data Rows
        for row_idx, row_data in enumerate(rows, start=2):
            ws.row_dimensions[row_idx].height = 20
            for col_idx, header in enumerate(headers, start=1):
                raw_val = row_data.get(header, "")
                cell_val = self._format_cell_value(raw_val)
                cell = ws.cell(row=row_idx, column=col_idx, value=cell_val)
                cell.font = self.body_font
                cell.border = self.cell_border

                # Alignment & Number Formatting
                if isinstance(cell_val, (int, float)) and not isinstance(cell_val, bool):
                    cell.alignment = self.right_align
                    if isinstance(cell_val, int):
                        cell.number_format = "#,##0"
                    else:
                        cell.number_format = "#,##0.00"
                else:
                    cell.alignment = self.left_align

        # 3. Auto-fit column widths
        for col in ws.columns:
            max_len = 0
            col_letter = get_column_letter(col[0].column)
            for cell in col:
                val_str = str(cell.value or "")
                # Korean characters take approx 2 width units
                width = sum(2 if ord(c) > 127 else 1 for c in val_str)
                if width > max_len:
                    max_len = width
            ws.column_dimensions[col_letter].width = max(max_len + 4, 12)

        wb.save(out_file)
        logger.info(f"Created styled Excel file: {out_file} ({len(rows)} rows)")
        return out_file
