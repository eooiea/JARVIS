"""Excel reader for parsing worksheets into structured Python dictionaries."""

from pathlib import Path
from typing import Optional, List, Dict, Any
import openpyxl
from jarvis.utils.logger import get_logger

logger = get_logger("office.excel_reader")


class ExcelReader:
    """Reads .xlsx files and converts sheet tables into row dictionaries."""

    def __init__(self):
        pass

    def read_excel(
        self,
        file_path: Path | str,
        sheet_name: Optional[str] = None,
    ) -> List[Dict[str, Any]]:
        """Read active or named worksheet into list of dictionaries."""
        path = Path(file_path).resolve()
        if not path.exists():
            raise FileNotFoundError(f"[엑셀 파일을 찾을 수 없습니다: {path}]")

        wb = openpyxl.load_workbook(path, data_only=True)
        if sheet_name and sheet_name in wb.sheetnames:
            ws = wb[sheet_name]
        else:
            ws = wb.active

        if ws is None or ws.max_row < 1:
            return []

        rows_iter = ws.iter_rows(values_only=True)
        header_row = next(rows_iter, None)
        if not header_row:
            return []

        headers = [str(h).strip() if h is not None else f"Column_{i+1}" for i, h in enumerate(header_row)]
        data_list = []

        for row in rows_iter:
            # Skip completely empty rows
            if not any(row):
                continue
            row_dict = {}
            for col_idx, header in enumerate(headers):
                val = row[col_idx] if col_idx < len(row) else None
                row_dict[header] = val
            data_list.append(row_dict)

        logger.info(f"Read {len(data_list)} rows from Excel: {path.name}")
        return data_list
