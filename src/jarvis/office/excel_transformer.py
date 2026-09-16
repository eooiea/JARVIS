"""High-performance Excel transformation and normalization engine for multi-channel orders."""

import re
from pathlib import Path
from typing import Dict, Any, List, Optional, Union
import openpyxl
from jarvis.utils.logger import get_logger

logger = get_logger("office.excel_transformer")

COURIER_PRESETS: Dict[str, Dict[str, Any]] = {
    "CJ대한통운": {
        "system": "eFLEXs / CNPlus",
        "target_schema": ["받는분성명", "전화번호1", "배송지주소", "품목명", "수량", "배송메시지"],
        "required_fields": ["받는분성명", "전화번호1", "배송지주소", "품목명", "수량"],
    },
    "우체국택배": {
        "system": "계약고객전용",
        "target_schema": ["수취인명", "수취인이동통신", "수취인주소", "내용품명", "수량", "배달메시지"],
        "required_fields": ["수취인명", "수취인이동통신", "수취인주소", "내용품명", "수량"],
    },
    "한진택배": {
        "system": "nFOCUS",
        "target_schema": ["받는사람", "받는분휴대폰", "주소", "상품명", "수량", "배송요청사항"],
        "required_fields": ["받는사람", "받는분휴대폰", "주소", "상품명", "수량"],
    },
    "롯데택배": {
        "system": "ALPS",
        "target_schema": ["수하인명", "수하인핸드폰", "수하인주소", "물품명", "수량", "전달메시지"],
        "required_fields": ["수하인명", "수하인핸드폰", "수하인주소", "물품명", "수량"],
    },
    "로젠택배": {
        "system": "iLogis",
        "target_schema": ["성명", "핸드폰", "배송지", "품목", "박스수량", "요청사항"],
        "required_fields": ["성명", "핸드폰", "배송지", "품목", "박스수량"],
    },
}


def normalize_phone_number(val: Any) -> str:
    """Normalize phone numbers into standard hyphenated format (e.g. 010-XXXX-XXXX)."""
    if val is None:
        return ""
    raw = str(val).strip()
    if not raw:
        return ""

    # Remove dots, spaces, hyphens
    digits = re.sub(r"[^\d]", "", raw)
    if not digits:
        return raw

    # 11 digits (010, 011, 016, etc. or 0504 safe numbers)
    if len(digits) == 11:
        if digits.startswith("050"):
            return f"{digits[:4]}-{digits[4:7]}-{digits[7:]}"
        return f"{digits[:3]}-{digits[3:7]}-{digits[7:]}"

    # 12 digits (0504-XXXX-XXXX)
    if len(digits) == 12 and digits.startswith("050"):
        return f"{digits[:4]}-{digits[4:8]}-{digits[8:]}"

    # 10 digits (02-XXXX-XXXX or 031-XXX-XXXX, 010-XXX-XXXX)
    if len(digits) == 10:
        if digits.startswith("02"):
            return f"{digits[:2]}-{digits[2:6]}-{digits[6:]}"
        return f"{digits[:3]}-{digits[3:6]}-{digits[6:]}"

    # 9 digits (02-XXX-XXXX)
    if len(digits) == 9 and digits.startswith("02"):
        return f"{digits[:2]}-{digits[2:5]}-{digits[5:]}"

    # Fallback to original
    return raw


def normalize_quantity(val: Any) -> int:
    """Extract integer quantity from various representations (e.g. '2개', '3EA', 4)."""
    if val is None or val == "":
        return 1
    if isinstance(val, int):
        return val
    try:
        return int(float(str(val)))
    except (ValueError, TypeError):
        pass

    match = re.search(r"\d+", str(val))
    if match:
        return int(match.group(0))
    return 1


def clean_address(val: Any) -> str:
    """Clean address string and collapse multiple spaces."""
    if val is None:
        return ""
    if isinstance(val, (list, tuple)):
        parts = [str(p).strip() for p in val if p is not None and str(p).strip()]
        combined = " ".join(parts)
        return re.sub(r"\s+", " ", combined).strip()
    
    cleaned = re.sub(r"\s+", " ", str(val)).strip()
    return cleaned


def get_courier_standard_preset(courier_name: str) -> Dict[str, Any]:
    """Retrieve standard schema preset for a major Korean courier service."""
    for name, preset in COURIER_PRESETS.items():
        if courier_name.strip().lower() in name.lower() or name.lower() in courier_name.strip().lower():
            return {
                "courier": name,
                "system": preset["system"],
                "target_schema": list(preset["target_schema"]),
                "required_fields": list(preset["required_fields"]),
            }
    
    # Defaults to CJ if unknown
    cj = COURIER_PRESETS["CJ대한통운"]
    return {
        "courier": courier_name,
        "system": "표준",
        "target_schema": list(cj["target_schema"]),
        "required_fields": list(cj["required_fields"]),
    }


def resolve_smart_path(input_path: Union[str, Path]) -> Path:
    """Resolve file path with fuzzy resilience against LLM typos and relative paths."""
    import difflib
    p = Path(input_path)
    if p.exists():
        return p.resolve()

    # Try resolving relative to CWD
    resolved_cwd = (Path.cwd() / p).resolve()
    if resolved_cwd.exists():
        return resolved_cwd

    # Try searching parent directory if parent exists
    parent = p.parent
    if not parent.exists():
        parent = (Path.cwd() / parent).resolve()

    if parent.exists() and parent.is_dir():
        target_name = p.name.lower()
        candidates = list(parent.glob("*.xlsx"))
        if candidates:
            cand_names = [c.name.lower() for c in candidates]
            # Exact or close match
            matches = difflib.get_close_matches(target_name, cand_names, n=1, cutoff=0.3)
            if matches:
                matched_idx = cand_names.index(matches[0])
                logger.info(f"Smart path resolved typo '{p.name}' -> '{candidates[matched_idx].name}'")
                return candidates[matched_idx]
            return candidates[0]

    return p.resolve()



class ExcelTransformer:
    """Performs schema inspection and multi-channel order transformation with zero hallucinations."""

    def __init__(self):
        pass

    def inspect(self, file_path: Union[str, Path], max_sample_rows: int = 1) -> Dict[str, Any]:
        """Inspect Excel sheet structure and headers with minimal token consumption."""
        path = resolve_smart_path(file_path)
        if not path.exists():
            return {
                "file_name": Path(file_path).name,
                "sheets": [],
                "headers": [],
                "total_rows": 0,
                "sample_rows": [],
                "error": f"엑셀 파일을 찾을 수 없습니다: {file_path}",
            }

        wb = openpyxl.load_workbook(path, data_only=True)
        sheet_names = wb.sheetnames
        ws = wb.active
        if ws is None or ws.max_row < 1:
            return {
                "file_name": path.name,
                "sheets": sheet_names,
                "headers": [],
                "total_rows": 0,
                "sample_rows": [],
            }


        rows_iter = ws.iter_rows(values_only=True)
        header_row = next(rows_iter, None)
        if not header_row:
            return {
                "file_name": path.name,
                "sheets": sheet_names,
                "headers": [],
                "total_rows": 0,
                "sample_rows": [],
            }

        headers = [str(h).strip() if h is not None else f"Column_{i+1}" for i, h in enumerate(header_row)]
        sample_rows = []
        total_rows = 0

        for row in rows_iter:
            if not any(row):
                continue
            total_rows += 1
            if len(sample_rows) < max_sample_rows:
                row_dict = {}
                for idx, h in enumerate(headers):
                    val = row[idx] if idx < len(row) else None
                    row_dict[h] = val
                sample_rows.append(row_dict)

        return {
            "file_name": path.name,
            "sheets": sheet_names,
            "headers": headers,
            "total_rows": total_rows,
            "sample_rows": sample_rows,
        }

    def transform(
        self,
        source_path: Union[str, Path],
        mapping_rules: Dict[str, Any],
        target_schema: Optional[List[str]] = None,
        required_fields: Optional[List[str]] = None,
    ) -> Dict[str, Any]:
        """Transform source Excel rows according to mapping rules into normalized target schema.
        
        Supports 1:1 mapping as well as hybrid multiple-source columns merged into a single target.
        """
        path = resolve_smart_path(source_path)
        if not path.exists():
            return {
                "success": False,
                "error_type": "FileNotFoundError",
                "message": f"파일을 찾을 수 없습니다: {source_path}",
            }

        # 1. Inspect source
        inspection = self.inspect(path, max_sample_rows=0)
        source_headers = inspection["headers"]

        # 2. Determine target schema and required fields
        actual_target_schema = target_schema or list(set(mapping_rules.values()))
        actual_required = required_fields or (
            [f for f in actual_target_schema if "메모" not in f and "메시지" not in f and "요청" not in f]
        )

        # 3. Invert/Index mapping rules: target_col -> list of source_col
        # Supports:
        # A) {"받는사람": "받는분성명", "기본배송지": "배송지주소", "상세배송지": "배송지주소"}
        # B) {"배송지주소": ["기본배송지", "상세배송지"]}
        target_to_sources: Dict[str, List[str]] = {}
        for k, v in mapping_rules.items():
            if isinstance(v, list):
                # Format B: target -> [source1, source2]
                target_col = k
                for src in v:
                    target_to_sources.setdefault(target_col, []).append(src)
            else:
                # Format A: source -> target
                target_col = str(v)
                src_col = str(k)
                target_to_sources.setdefault(target_col, []).append(src_col)

        # 4. Strict Schema Validation (Gatekeeper)
        missing_required = []
        for req in actual_required:
            if req not in target_to_sources or not target_to_sources[req]:
                missing_required.append(req)

        if missing_required:
            return {
                "success": False,
                "error_type": "SchemaValidationError",
                "missing_fields": missing_required,
                "message": f"타깃 스키마의 필수 필드 매핑이 누락되었습니다: {missing_required}",
                "hint": f"mapping_rules에 누락된 필드 {missing_required}에 대응하는 소스 엑셀 컬럼을 추가 지정하여 다시 호출하십시오.",
            }

        # 5. Read all rows and transform
        wb = openpyxl.load_workbook(path, data_only=True)
        ws = wb.active
        if ws is None or ws.max_row < 2:
            return {"success": True, "count": 0, "rows": []}

        rows_iter = ws.iter_rows(values_only=True)
        header_row = next(rows_iter, None)
        headers = [str(h).strip() if h is not None else f"Column_{i+1}" for i, h in enumerate(header_row)]

        transformed_rows: List[Dict[str, Any]] = []

        for row in rows_iter:
            if not any(row):
                continue

            src_row_dict = {h: (row[i] if i < len(row) else None) for i, h in enumerate(headers)}
            new_row: Dict[str, Any] = {}

            # Construct fields for target_schema
            for t_col in actual_target_schema:
                src_cols = target_to_sources.get(t_col, [])
                if not src_cols:
                    new_row[t_col] = ""
                    continue

                if len(src_cols) == 1:
                    raw_val = src_row_dict.get(src_cols[0])
                else:
                    # Multi-column merge (e.g. split address or name + option)
                    vals = [str(src_row_dict.get(sc, "")).strip() for sc in src_cols if src_row_dict.get(sc) is not None]
                    raw_val = " ".join(filter(None, vals))

                # Normalize by column semantics
                col_name_lower = t_col.lower()
                if any(kw in col_name_lower for kw in ("전화번호", "휴대폰", "핸드폰", "연락처", "통신")):
                    new_row[t_col] = normalize_phone_number(raw_val)
                elif any(kw in col_name_lower for kw in ("수량", "박스수량", "개수", "qty")):
                    new_row[t_col] = normalize_quantity(raw_val)
                elif any(kw in col_name_lower for kw in ("주소", "배송지", "배달지")):
                    new_row[t_col] = clean_address(raw_val)
                else:
                    new_row[t_col] = str(raw_val).strip() if raw_val is not None else ""

            transformed_rows.append(new_row)

        logger.info(f"Successfully transformed {len(transformed_rows)} rows from {path.name}")
        return {
            "success": True,
            "count": len(transformed_rows),
            "rows": transformed_rows,
        }


def auto_map_headers(source_headers: List[str], target_schema: List[str]) -> Dict[str, Any]:
    """Intelligently map source headers to target schema based on Korean ecommerce semantics."""
    mapping_rules: Dict[str, Any] = {}

    # Order matters: Specific keywords (phone, address, quantity) MUST be checked before recipient
    # so that '수하인핸드폰' is classified as phone, not recipient!
    semantic_priorities = [
        ("phone", ["전화번호", "휴대폰", "핸드폰", "연락처", "이동통신", "휴대전화", "수신자연락처"]),
        ("address", ["주소", "배송지", "배달지", "기본배송지", "상세배송지", "수령주소", "전체주소", "주소지"]),
        ("quantity", ["수량", "구매수량", "주문수량", "출고수량", "박스수량", "개수", "qty"]),
        ("memo", ["메시지", "메세지", "요구사항", "요청사항", "메모", "배송특이사항", "요청", "배송요청", "배송메모"]),
        ("product", ["상품명", "품목", "내용품", "물품", "등록상품명", "아이템", "출고상품", "주문상품명", "상품", "상품상세", "아이템명"]),
        ("recipient", ["성명", "받는사람", "받는분", "수취인", "수하인", "이름", "고객명", "수령인", "수령자", "구매자명", "구매자", "받는사람이름", "받는분성함"]),
    ]


    # Identify semantic category of each target column
    for t_col in target_schema:
        t_lower = t_col.lower()
        matched_category = None
        for cat, keywords in semantic_priorities:
            if any(kw in t_lower for kw in keywords):
                matched_category = cat
                break

        if not matched_category:
            # Fallback: direct match
            for s_col in source_headers:
                if s_col == t_col:
                    mapping_rules[s_col] = t_col
            continue

        # Find matching source headers
        matched_sources = []
        for s_col in source_headers:
            s_lower = s_col.lower()
            # Find category for source header
            src_cat = None
            for cat, keywords in semantic_priorities:
                if any(kw in s_lower for kw in keywords):
                    src_cat = cat
                    break
            
            if src_cat == matched_category:
                # Avoid buyer phone if recipient phone exists
                if matched_category == "phone" and "구매자" in s_lower and any("수하인" in x or "수취인" in x for x in source_headers):
                    continue
                matched_sources.append(s_col)

        if matched_category == "address" and len(matched_sources) > 1:
            # Multi-column split address (e.g. 기본배송지, 상세배송지)
            for src in matched_sources:
                mapping_rules[src] = t_col
        elif matched_sources:
            mapping_rules[matched_sources[0]] = t_col

    return mapping_rules



def batch_transform_orders(
    order_dir: Union[str, Path],
    courier_name: str,
    output_path: Union[str, Path],
    template_path: Optional[Union[str, Path]] = None,
) -> Dict[str, Any]:
    """Batch collect multi-channel order Excel files from directory, normalize, and generate invoice Excel."""
    from jarvis.office.excel_writer import ExcelWriter

    dir_path = resolve_smart_path(order_dir)
    if not dir_path.exists() or not dir_path.is_dir():
        return {
            "success": False,
            "error": f"주문서 폴더를 찾을 수 없습니다: {order_dir}",
        }

    # Find all Excel files
    excel_files = sorted(list(dir_path.glob("*.xlsx")))
    if not excel_files:
        return {
            "success": False,
            "error": f"폴더 내에 .xlsx 주문서 파일이 없습니다: {dir_path}",
        }

    transformer = ExcelTransformer()

    # Determine target schema
    target_schema = []
    required_fields = []

    if template_path:
        tpl_path = resolve_smart_path(template_path)
        if tpl_path.exists() and tpl_path.is_file():
            tpl_info = transformer.inspect(tpl_path, max_sample_rows=0)
            target_schema = tpl_info.get("headers", [])
            if target_schema:
                required_fields = [h for h in target_schema if "메모" not in h and "특이" not in h]

    if not target_schema:
        preset = get_courier_standard_preset(courier_name)
        target_schema = preset["target_schema"]
        required_fields = preset["required_fields"]


    all_transformed_rows: List[Dict[str, Any]] = []
    processed_files: List[str] = []

    for f in excel_files:
        inspection = transformer.inspect(f, max_sample_rows=0)
        source_headers = inspection.get("headers", [])
        rules = auto_map_headers(source_headers, target_schema)

        res = transformer.transform(
            source_path=f,
            mapping_rules=rules,
            target_schema=target_schema,
            required_fields=required_fields,
        )
        if res.get("success"):
            all_transformed_rows.extend(res.get("rows", []))
            processed_files.append(f.name)
        else:
            logger.warning(f"Failed to transform file {f.name}: {res.get('message')}")

    out_p = Path(output_path).resolve()
    out_p.parent.mkdir(parents=True, exist_ok=True)

    writer = ExcelWriter()
    created = writer.create_styled_excel(
        data=all_transformed_rows,
        output_path=out_p,
        sheet_name=f"{courier_name}송장",
    )

    return {
        "success": True,
        "courier": courier_name,
        "processed_files": processed_files,
        "total_orders": len(all_transformed_rows),
        "output_path": str(created),
        "target_schema": target_schema,
    }

