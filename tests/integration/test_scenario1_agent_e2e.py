"""End-to-End Autonomous Agent Execution & Log Observation for Scenario 1 (10 Test Sets)."""

import re
import json
from pathlib import Path
import pytest
import openpyxl

from jarvis.main import JarvisPipeline
from jarvis.utils.config import load_config
from jarvis.utils.logger import get_logger

logger = get_logger("test.scenario1_e2e")


def analyze_session_log(log_path: str) -> dict:
    """Read and evaluate agent session execution log without dumping raw contents."""
    p = Path(log_path)
    if not p.exists():
        return {"error": "Log file not found"}

    content = p.read_text(encoding="utf-8")
    turns = len(re.findall(r"## 🔄 \[Turn (\d+)\]", content))
    thoughts = re.findall(r"\*\*Thought\*\*:\s*([^\n]+)", content)
    tool_calls = re.findall(r"### 🛠️ \[도구 호출: ([\w_]+)\]", content)
    has_completion = "## 🏁 [최종 완결 (Task Completed)]" in content

    return {
        "log_file": p.name,
        "turns_used": turns,
        "thought_count": len(thoughts),
        "tool_call_count": len(tool_calls),
        "tool_names": tool_calls,
        "completed": has_completion,
    }


def verify_generated_invoice(
    excel_path: Path,
    expected_rows_count: int,
    expected_headers: list,
    phone_col_idx: int = 1,
):
    """Verify generated invoice Excel file for 1:1 row conservation and phone normalization."""
    target = excel_path
    if not target.exists():
        candidates = list(target.parent.glob("*.xlsx"))
        if candidates:
            target = candidates[0]

    assert target.exists(), f"결과 송장 파일이 생성되지 않았습니다: {excel_path}"
    wb = openpyxl.load_workbook(target)
    ws = wb.active
    assert ws is not None

    # Verify headers
    headers = [cell.value for cell in ws[1]]
    for h in expected_headers:
        assert h in headers, f"필수 헤더 '{h}'가 결과 엑셀에 없습니다. 실제 헤더: {headers}"

    # Verify row count (1:1 preservation: header row + data rows)
    actual_rows = ws.max_row - 1
    assert actual_rows == expected_rows_count, (
        f"주문 건수 1:1 보존 실패! 기대: {expected_rows_count}건, 실제: {actual_rows}건"
    )

    # Verify phone number normalization in data rows
    for r in range(2, ws.max_row + 1):
        phone_val = str(ws.cell(row=r, column=phone_col_idx + 1).value or "")
        if phone_val:
            assert "." not in phone_val, f"전화번호에 온점이 남아있습니다: {phone_val}"
            assert "-" in phone_val or len(phone_val) == 0, f"전화번호 하이픈 정규화 누락: {phone_val}"


class TestScenario1AllSets:
    """Execute and verify JARVIS autonomous performance across all 10 scenario sets."""

    @pytest.fixture(autouse=True)
    def setup(self):
        self.config = load_config()
        self.base_dir = Path("tests/data/scenario1").resolve()
        self.log_dir = Path("logs").resolve()
        self.log_dir.mkdir(parents=True, exist_ok=True)

    def test_set01_basic_cj_standard(self, tmp_path: Path):
        """Set 01: Coupang + Smartstore + Gmarket -> CJ (11 rows total, 1:1 preserved)."""
        target_dir = self.base_dir / "set01" / "orders"
        out_excel = tmp_path / "CJ대한통운_송장.xlsx"

        pipeline = JarvisPipeline(config=self.config, output_dir=tmp_path, log_dir=self.log_dir)
        instruction = (
            f"'{target_dir}' 폴더의 주문 엑셀 파일들을 모두 읽어서 CJ대한통운 표준 송장 양식으로 정규화하여 "
            f"'{out_excel}' 파일로 일괄 변환해줘."
        )

        res = pipeline.execute_instruction(instruction=instruction, auto_confirm=True)
        assert res["success"] is True

        verify_generated_invoice(
            excel_path=out_excel,
            expected_rows_count=11,
            expected_headers=["받는분성명", "전화번호1", "배송지주소", "품목명", "수량"],
            phone_col_idx=1,
        )

    def test_set02_epost_standard(self, tmp_path: Path):
        """Set 02: Smartstore + 11st -> Epost (4 rows total)."""
        target_dir = self.base_dir / "set02" / "orders"
        out_excel = tmp_path / "우체국택배_송장.xlsx"

        pipeline = JarvisPipeline(config=self.config, output_dir=tmp_path, log_dir=self.log_dir)
        instruction = (
            f"'{target_dir}' 안의 엑셀 주문서들을 읽어서 우체국택배 표준 양식으로 변환해서 '{out_excel}'로 저장해줘."
        )

        res = pipeline.execute_instruction(instruction=instruction, auto_confirm=True)
        assert res["success"] is True

        verify_generated_invoice(
            excel_path=out_excel,
            expected_rows_count=4,
            expected_headers=["수취인명", "수취인이동통신", "수취인주소", "내용품명", "수량"],
            phone_col_idx=1,
        )

    def test_set03_hanjin_standard(self, tmp_path: Path):
        """Set 03: Coupang + Kakao -> Hanjin (3 rows total)."""
        target_dir = self.base_dir / "set03" / "orders"
        out_excel = tmp_path / "한진택배_송장.xlsx"

        pipeline = JarvisPipeline(config=self.config, output_dir=tmp_path, log_dir=self.log_dir)
        instruction = f"'{target_dir}' 주문 목록을 한진택배 송장 양식으로 변환하여 '{out_excel}'로 생성해줘."
        res = pipeline.execute_instruction(instruction=instruction, auto_confirm=True)
        assert res["success"] is True

        verify_generated_invoice(
            excel_path=out_excel,
            expected_rows_count=3,
            expected_headers=["받는사람", "받는분휴대폰", "주소", "상품명", "수량"],
            phone_col_idx=1,
        )

    def test_set04_lotte_standard(self, tmp_path: Path):
        """Set 04: Musinsa + 29CM -> Lotte (3 rows total)."""
        target_dir = self.base_dir / "set04" / "orders"
        out_excel = tmp_path / "롯데택배_송장.xlsx"

        pipeline = JarvisPipeline(config=self.config, output_dir=tmp_path, log_dir=self.log_dir)
        instruction = f"'{target_dir}' 의 패션 주문서들을 롯데택배 양식으로 일괄 변환하여 '{out_excel}'에 저장해줘."
        res = pipeline.execute_instruction(instruction=instruction, auto_confirm=True)
        assert res["success"] is True

        verify_generated_invoice(
            excel_path=out_excel,
            expected_rows_count=3,
            expected_headers=["수하인명", "수하인핸드폰", "수하인주소", "물품명", "수량"],
            phone_col_idx=1,
        )

    def test_set05_logen_standard(self, tmp_path: Path):
        """Set 05: Gmarket + Auction -> Logen (3 rows total)."""
        target_dir = self.base_dir / "set05" / "orders"
        out_excel = tmp_path / "로젠택배_송장.xlsx"

        pipeline = JarvisPipeline(config=self.config, output_dir=tmp_path, log_dir=self.log_dir)
        instruction = f"'{target_dir}' 내 ESM 주문 파일들을 로젠택배 송장 규격으로 변환해서 '{out_excel}'로 만들어줘."
        res = pipeline.execute_instruction(instruction=instruction, auto_confirm=True)
        assert res["success"] is True

        verify_generated_invoice(
            excel_path=out_excel,
            expected_rows_count=3,
            expected_headers=["성명", "핸드폰", "배송지", "품목", "박스수량"],
            phone_col_idx=1,
        )

    def test_set06_custom_template(self, tmp_path: Path):
        """Set 06: Smartstore + Coupang -> Custom User Template (2 rows total)."""
        target_dir = self.base_dir / "set06" / "orders"
        template_file = self.base_dir / "set06" / "templates" / "자사_송장양식.xlsx"
        out_excel = tmp_path / "커스텀_자사송장.xlsx"

        pipeline = JarvisPipeline(config=self.config, output_dir=tmp_path, log_dir=self.log_dir)
        instruction = (
            f"'{template_file}' 양식의 헤더에 맞춰서 '{target_dir}' 주문 파일들을 취합 정리하여 "
            f"'{out_excel}' 파일로 만들어줘."
        )
        res = pipeline.execute_instruction(instruction=instruction, auto_confirm=True)
        assert res["success"] is True

        verify_generated_invoice(
            excel_path=out_excel,
            expected_rows_count=2,
            expected_headers=["고객명", "휴대폰번호", "전체주소", "출고상품", "출고수량"],
            phone_col_idx=1,
        )

    def test_set07_fashion_beauty_cj(self, tmp_path: Path):
        """Set 07: Ably + Zigzag + OliveYoung -> CJ (3 rows total)."""
        target_dir = self.base_dir / "set07" / "orders"
        out_excel = tmp_path / "CJ_뷰티_송장.xlsx"

        pipeline = JarvisPipeline(config=self.config, output_dir=tmp_path, log_dir=self.log_dir)
        instruction = f"'{target_dir}' 주문서들을 취합해서 CJ대한통운 양식의 '{out_excel}'로 변환해줘."
        res = pipeline.execute_instruction(instruction=instruction, auto_confirm=True)
        assert res["success"] is True

        verify_generated_invoice(
            excel_path=out_excel,
            expected_rows_count=3,
            expected_headers=["받는분성명", "전화번호1", "배송지주소", "품목명", "수량"],
            phone_col_idx=1,
        )

    def test_set08_phone_noise_epost(self, tmp_path: Path):
        """Set 08: Safe number and landline numbers -> Epost (3 rows total)."""
        target_dir = self.base_dir / "set08" / "orders"
        out_excel = tmp_path / "우체국_전화번호정제_송장.xlsx"

        pipeline = JarvisPipeline(config=self.config, output_dir=tmp_path, log_dir=self.log_dir)
        instruction = f"'{target_dir}' 주문서들을 우체국택배 양식으로 변환해서 '{out_excel}'로 저장해줘."
        res = pipeline.execute_instruction(instruction=instruction, auto_confirm=True)
        assert res["success"] is True

        verify_generated_invoice(
            excel_path=out_excel,
            expected_rows_count=3,
            expected_headers=["수취인명", "수취인이동통신", "수취인주소", "내용품명"],
            phone_col_idx=1,
        )

    def test_set09_bulk_orders_lotte(self, tmp_path: Path):
        """Set 09: Bulk 15 orders -> Lotte (15 rows total 1:1 preserved)."""
        target_dir = self.base_dir / "set09" / "orders"
        out_excel = tmp_path / "롯데_대량송장.xlsx"

        pipeline = JarvisPipeline(config=self.config, output_dir=tmp_path, log_dir=self.log_dir)
        instruction = f"'{target_dir}' 대량 발주 주문서를 롯데택배 규격으로 변환해서 '{out_excel}'로 생성해줘."
        res = pipeline.execute_instruction(instruction=instruction, auto_confirm=True)
        assert res["success"] is True

        verify_generated_invoice(
            excel_path=out_excel,
            expected_rows_count=15,
            expected_headers=["수하인명", "수하인핸드폰", "수하인주소", "물품명", "수량"],
            phone_col_idx=1,
        )

    def test_set10_self_correction_logen(self, tmp_path: Path):
        """Set 10: 3 platforms -> Logen with self-correction feedback loop (3 rows total)."""
        target_dir = self.base_dir / "set10" / "orders"
        out_excel = tmp_path / "로젠_자가수정_송장.xlsx"

        pipeline = JarvisPipeline(config=self.config, output_dir=tmp_path, log_dir=self.log_dir)
        instruction = f"'{target_dir}' 안의 파일들을 로젠택배 규격으로 완벽히 매핑하여 '{out_excel}'로 저장해줘."
        res = pipeline.execute_instruction(instruction=instruction, auto_confirm=True)
        assert res["success"] is True

        verify_generated_invoice(
            excel_path=out_excel,
            expected_rows_count=3,
            expected_headers=["성명", "핸드폰", "배송지", "품목", "박스수량"],
            phone_col_idx=1,
        )
