"""Unit tests for JARVIS Smart Office Tools (TDD Red-Green)."""

import openpyxl
import pytest
from pathlib import Path
from jarvis.tools.office import create_excel


def test_create_excel_from_list_of_rows(tmp_path: Path):
    """create_excel should generate styled workbook from list of dictionaries."""
    out_file = tmp_path / "folder_structure.xlsx"
    data = [
        {"이름": "src", "종류": "폴더", "크기": "4 KB"},
        {"이름": "main.py", "종류": "파일", "크기": "2.5 KB"},
        {"이름": "config.yaml", "종류": "파일", "크기": "1.1 KB"},
    ]

    result_path = create_excel(data, output_path=out_file)
    assert result_path.exists()

    wb = openpyxl.load_workbook(result_path)
    ws = wb.active
    assert ws is not None
    # Header check
    assert ws["A1"].value == "이름"
    assert ws["B1"].value == "종류"
    # Row 1 check
    assert ws["A2"].value == "src"
    assert ws["B2"].value == "폴더"
    # Row 2 check
    assert ws["A3"].value == "main.py"
    # Total row count
    assert ws.max_row == 4
