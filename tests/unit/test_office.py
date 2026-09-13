"""Unit tests for JARVIS Office & Smart Excel Engine (TDD Red-Green)."""

import openpyxl
import pytest
from pathlib import Path
from unittest.mock import MagicMock

from jarvis.office.file_manager import FileManager
from jarvis.office.file_namer import FileNamer
from jarvis.office.excel_writer import ExcelWriter
from jarvis.office.excel_reader import ExcelReader
from jarvis.office.csv_handler import CSVHandler


# ==========================================
# 1. FileManager Tests (Duplicate Handling)
# ==========================================
def test_file_manager_unique_filepath(tmp_path: Path):
    """FileManager should append incrementing numbers to prevent overwriting existing files."""
    manager = FileManager(output_dir=tmp_path)
    
    # 1st file: report.xlsx
    file1 = manager.get_unique_filepath("report.xlsx")
    assert file1.name == "report.xlsx"
    file1.write_text("file1 content")

    # 2nd file: report_(1).xlsx
    file2 = manager.get_unique_filepath("report.xlsx")
    assert file2.name == "report_(1).xlsx"
    file2.write_text("file2 content")

    # 3rd file: report_(2).xlsx
    file3 = manager.get_unique_filepath("report.xlsx")
    assert file3.name == "report_(2).xlsx"


# ==========================================
# 2. FileNamer Tests
# ==========================================
def test_file_namer_generates_formatted_name():
    """FileNamer should generate date-prefixed meaningful filename from document data."""
    namer = FileNamer()
    data = {"title": "회식 영수증", "date": "2026-09-02"}
    filename = namer.generate_filename(data, default_prefix="영수증_정리")
    
    assert filename.endswith(".xlsx")
    assert "회식_영수증" in filename or "영수증" in filename


# ==========================================
# 3. ExcelWriter & ExcelReader Tests
# ==========================================
def test_excel_writer_creates_styled_sheet(tmp_path: Path):
    """ExcelWriter should create styled .xlsx with header color, borders, and proper formatting."""
    writer = ExcelWriter()
    output_file = tmp_path / "styled_output.xlsx"

    sample_data = {
        "title": "비용 정산서",
        "items": [
            {"품목": "복사용지", "수량": 5, "단가": 3000, "금액": 15000},
            {"품목": "토너 카트리지", "수량": 1, "단가": 85000, "금액": 85000}
        ]
    }

    result_path = writer.create_styled_excel(sample_data, output_path=output_file)
    assert result_path.exists()

    # Verify formatting with openpyxl
    wb = openpyxl.load_workbook(result_path)
    ws = wb.active
    assert ws is not None

    # Check header row (Bold font, fill color)
    header_cell = ws["A1"]
    assert header_cell.value == "품목"
    assert header_cell.font.bold is True
    assert header_cell.fill.fill_type == "solid"

    # Verify data content
    assert ws["A2"].value == "복사용지"
    assert ws["D2"].value == 15000
    assert ws["A3"].value == "토너 카트리지"
    assert ws["D3"].value == 85000


def test_excel_writer_handles_arbitrary_nested_dicts(tmp_path: Path):
    """ExcelWriter should safely convert arbitrary nested dictionaries and lists without crashing."""
    writer = ExcelWriter()
    output_file = tmp_path / "nested_output.xlsx"

    nested_data = {
        "document_name": "프로젝트 문서",
        "sections": [
            {"section_number": "1", "title": "개요 및 목표", "content": "시스템 목표입니다.", "metadata": {"author": "admin"}},
            {"section_number": "2", "title": "세부 구현", "content": "상세 내용입니다.", "metadata": {"author": "dev"}}
        ]
    }

    result_path = writer.create_styled_excel(nested_data, output_path=output_file)
    assert result_path.exists()

    wb = openpyxl.load_workbook(result_path)
    ws = wb.active
    assert ws is not None
    assert ws["A1"].value == "section_number"
    assert ws["A2"].value == "1"
    # metadata dict is safely converted to string
    assert "author" in str(ws["D2"].value)


def test_excel_reader_extracts_data(tmp_path: Path):
    """ExcelReader should read existing excel file and parse rows into list of dictionaries."""
    writer = ExcelWriter()
    excel_path = tmp_path / "sample_read.xlsx"
    data = [
        {"이름": "홍길동", "부서": "개발팀", "직급": "선임"},
        {"이름": "김철수", "부서": "기획팀", "직급": "책임"}
    ]
    writer.create_styled_excel(data, output_path=excel_path)

    reader = ExcelReader()
    rows = reader.read_excel(excel_path)

    assert len(rows) == 2
    assert rows[0]["이름"] == "홍길동"
    assert rows[0]["부서"] == "개발팀"
    assert rows[1]["이름"] == "김철수"


# ==========================================
# 4. CSVHandler Tests
# ==========================================
def test_csv_handler_roundtrip(tmp_path: Path):
    """CSVHandler should export UTF-8-SIG CSV and read it back without Korean encoding issues."""
    handler = CSVHandler()
    csv_file = tmp_path / "data.csv"
    data = [
        {"날짜": "2026-09-02", "적요": "식대 결제", "금액": "12000"},
        {"날짜": "2026-09-03", "적요": "사무용품", "금액": "35000"}
    ]

    saved_path = handler.export_csv(data, output_path=csv_file)
    assert saved_path.exists()

    loaded = handler.read_csv(saved_path)
    assert len(loaded) == 2
    assert loaded[0]["적요"] == "식대 결제"
    assert loaded[1]["금액"] == "35000"
